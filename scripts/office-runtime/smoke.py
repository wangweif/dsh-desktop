"""Exercise installed authoring libraries and OPC paths, using only bundled Python."""
import importlib.metadata
import json
import pathlib
import sys
import zipfile
from docx import Document
from pptx import Presentation
from openpyxl import Workbook, load_workbook
import numpy
import pandas
import PIL.Image
import lxml.etree
import xlsxwriter

root = pathlib.Path(sys.argv[1])
manifest = json.loads(pathlib.Path(sys.argv[2]).read_text(encoding='utf-8'))
assert '.'.join(map(str, sys.version_info[:3])) == manifest['python']
for name, version in manifest['pythonPackages'].items():
    assert importlib.metadata.version(name) == version, name
assert numpy.arange(4).sum() == 6
assert pandas.DataFrame({'n': [1, 2]}).n.sum() == 3
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading('Office runtime smoke', 0)
doc.add_paragraph('Bundled Python authoring')
doc.save(root / 'sample.docx')
assert Document(root / 'sample.docx').paragraphs[1].text == 'Bundled Python authoring'
ppt = Presentation()
ppt.slides.add_slide(ppt.slide_layouts[0]).shapes.title.text = 'Office runtime smoke'
ppt.save(root / 'sample.pptx')
assert Presentation(root / 'sample.pptx').slides[0].shapes.title.text == 'Office runtime smoke'
book = Workbook()
book.active['A1'] = 'Office runtime smoke'
book.active['A2'] = 7
book.active['A3'] = '=A2*2'
book.save(root / 'sample.xlsx')
assert load_workbook(root / 'sample.xlsx').active['A3'].value == '=A2*2'
writer = xlsxwriter.Workbook(root / 'writer.xlsx')
writer.add_worksheet().write('A1', 'Office runtime smoke')
writer.close()
for file, part in [('sample.docx', 'word/document.xml'), ('sample.pptx', 'ppt/presentation.xml'), ('sample.xlsx', 'xl/workbook.xml'), ('writer.xlsx', 'xl/workbook.xml')]:
    with zipfile.ZipFile(root / file) as archive:
        assert part in archive.namelist()
        assert '_rels/.rels' in archive.namelist()
        assert not any('\\' in name for name in archive.namelist()), file
print('Bundled Python created and reopened DOCX/PPTX/XLSX with OPC slash paths')
